"""Independent final checks against selected IDs and all stored Dev predictions."""
import hashlib
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import accuracy_score,f1_score

import run_cross_verifier as ex


def main():
    manifest=ex.check_frozen()
    results=pd.read_csv(ex.OUT/'cross_verifier_run_results.csv')
    assert len(results)==60
    assert results.groupby('Method').size().to_dict()=={'Random320':25,**{x:5 for x in ex.MAIN_METHODS[1:]+ex.SOURCE_METHODS}}
    dev=pd.read_csv(ex.DEV)
    audit=json.loads((ex.OUT/'evaluation_audit.json').read_text())
    assert audit['dev_sha256']==ex.old.digest(ex.DEV)
    assert not audit['final_holdout_accessed']
    scores=pd.read_csv(ex.OUT/'candidate_verifier_scores.csv',float_precision='round_trip')
    pool=pd.DataFrame(ex.pool_rows())
    for path,digest in manifest['selected_dataset_hashes'].items():
        assert ex.old.digest(ex.OUT/path)==digest
        selected=pd.DataFrame(ex.old.read_jsonl(ex.OUT/path))
        assert set(selected.groupby(['source','label']).size())=={40}
        name=Path(path).stem
        if name in ex.MAIN_METHODS[1:]:
            expected=set()
            for (source,label),g in scores.groupby(['source','label']):
                if name=='TopK320':column='proximity_score'
                else:
                    verifier=source if name=='SelfVerifier320' else {'qwen':'llama','llama':'qwen'}[source]
                    column=verifier+'_LogLabelScore'
                expected.update(g.sort_values([column,'sample_id'],ascending=[False,True]).head(40).sample_id)
            assert set(selected.sample_id)==expected
    for _,row in results.iterrows():
        prediction=pd.read_csv(Path(row.Run_Dir)/'dev_predictions.csv')
        assert len(prediction)==2000
        assert prediction.sample_id.tolist()==dev.sample_id.tolist()
        assert np.array_equal(prediction.true_label,dev.label)
        expected=f1_score(prediction.true_label,prediction.predicted_label,labels=[0,1,2,3],average=None,zero_division=0)
        assert abs(accuracy_score(prediction.true_label,prediction.predicted_label)-row.Dev_Accuracy)<1e-12
        assert abs(expected.mean()-row.Dev_Macro_F1)<1e-12
        for label,value in zip(ex.repair.LABELS,expected):assert abs(value-row[f'Dev_{label}_F1'])<1e-12
        assert row.training_steps==(30 if row.Method in ex.SOURCE_METHODS else 60)
        assert all(math.isfinite(row[x]) for x in ex.METRICS)
    assert all(g.classifier_init_sha256.nunique()==1 for _,g in results.groupby('Model_Seed'))
    random=results[results.Method=='Random320'].groupby('Model_Seed')[ex.METRICS].mean()
    paired=pd.read_csv(ex.OUT/'paired_model_seed_deltas.csv')
    assert len(paired)==35
    for _,row in paired.iterrows():
        a=results[(results.Method==row.Left)&(results.Model_Seed==row.Model_Seed)].iloc[0]
        b=random.loc[row.Model_Seed] if row.Right=='RandomMean' else results[(results.Method==row.Right)&(results.Model_Seed==row.Model_Seed)].iloc[0]
        for metric in ex.METRICS:assert abs(row[metric+'_Delta']-(a[metric]-b[metric]))<1e-12
    workbook=load_workbook(ex.OUT/'cross_verifier_summary.xlsx',read_only=True)
    assert workbook['RunResults'].max_row==61
    assert workbook['PairedSeedDeltas'].max_row==36
    for model in ['qwen','llama']:
        records=ex.repair.cached(ex.OUT/f'{model}_scores.jsonl')
        assert set(records)==set(pool.sample_id)
    evidence=dict(passed=True,run_count=60,dev_prediction_rows_recomputed=120000,selection_hashes_checked=12,complete_verifier_scores=8000,classifier_seed_pairing=True,paired_delta_rows_checked=35,final_holdout_accessed=False)
    ex.old.write_json(ex.OUT/'final_validation.json',evidence)
    print(json.dumps(evidence,indent=2))


if __name__=='__main__':main()
