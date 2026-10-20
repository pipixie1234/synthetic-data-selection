#!/usr/bin/env python3
"""Build faculty-discussion figures from existing workbooks; never run experiments.

Run from any directory: python tools/build_research_presentation.py
Dependencies: python -m pip install -r requirements-presentation.txt
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from PIL import Image
import xlsxwriter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "成果"
FIG = OUT / "PPT图表"
BOOKS = OUT / "关键Excel"
TABLES = OUT / "图表数据"
SOURCES = {
    "topk": ("01_TopK多seed验证.xlsx", "SyntheticData/exp/outputs/curation_validation/e1_e2_validation_results.xlsx"),
    "matrix": ("02_来源兼容性矩阵.xlsx", "SyntheticData/SmaExp/outputs/visualizations/comparison_summary_500_matrix.xlsx"),
    "ln": ("03_LN_DSIR与ESS诊断.xlsx", "SyntheticData/DisrTry/outputs/ln_dsir_validation/ln_dsir_results.xlsx"),
    "car": ("04_CaR筛选器对照.xlsx", "SyntheticData/CaRTry/outputs/car_selector_validation/car_selector_results.xlsx"),
    "verifier": ("05_Verifier下游验证.xlsx", "SyntheticData/cross_verifier/outputs/cross_verifier_summary.xlsx"),
    "dsir": ("06_原始DSIR对照.xlsx", "SyntheticData/DisrTry/outputs/dsir_selector_validation/dsir_selector_results.xlsx"),
}
NAVY, TEAL, GRAY, AMBER = "#15314B", "#008C95", "#8C9BAB", "#DB9145"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 13,
                     "axes.spines.top": False, "axes.spines.right": False,
                     "axes.labelcolor": NAVY, "text.color": NAVY,
                     "axes.titleweight": "bold", "svg.fonttype": "none"})


def sheet(key, name):
    w = load_workbook(ROOT / SOURCES[key][1], data_only=True, read_only=True)
    rows = list(w[name].values)
    w.close()
    return [dict(zip(rows[0], r)) for r in rows[1:] if any(v is not None for v in r)]


def savefig(fig, name, foot):
    fig.text(.06, .025, foot, fontsize=10, color="#53697D", va="bottom")
    fig.subplots_adjust(left=.09, right=.96, bottom=.19, top=.84)
    fig.savefig(FIG / (name + ".png"), dpi=300, facecolor="white")
    fig.savefig(FIG / (name + ".svg"), facecolor="white")
    plt.close(fig)


def main():
    for d in (OUT, FIG, BOOKS, TABLES):
        d.mkdir(parents=True, exist_ok=True)
    manifest = []
    for key, (name, rel) in SOURCES.items():
        original = ROOT / rel
        target = BOOKS / name
        shutil.copy2(original, target)
        sha = hashlib.sha256(original.read_bytes()).hexdigest()
        assert hashlib.sha256(target.read_bytes()).hexdigest() == sha
        manifest.append({"key": key, "copy": str(target.relative_to(OUT)), "source": rel, "sha256": sha})

    paired = sheet("topk", "Model_Seed_Results")
    classes = sheet("topk", "Per_Class_F1")
    matrix_rows = sheet("matrix", "Accuracy")
    ln = sheet("ln", "Main_Comparison")
    car = sheet("car", "Main_Comparison")
    ess = sheet("ln", "ESS_Comparison")
    verifier = sheet("verifier", "MethodSummary")
    lookup = {(r["Experiment"], r["Model_Seed"], r["Metric"]): r["Mean"] for r in paired}
    seeds = [42, 43, 44, 45, 46]
    data = {}
    panels = []
    fig, axes = plt.subplots(1, 2, figsize=(12, 5.7))
    for ax, metric, label in zip(axes, ["Dev_Accuracy", "Dev_Macro_F1"], ["Accuracy", "Macro-F1"]):
        vals = np.array([[lookup[(method, seed, metric)] * 100 for method in ("E1", "E2")] for seed in seeds])
        panels.append(vals)
        for seed, values in zip(seeds, vals):
            ax.plot([0, 1], values, "o-", color=GRAY, alpha=.75, lw=1.5, ms=6)
            ax.text(1.03, values[1], str(seed), fontsize=9, color=GRAY, va="center")
        means = vals.mean(axis=0)
        ax.plot([0, 1], means, "o-", color=TEAL, lw=3.4, ms=10, label="Mean")
        for x, y in enumerate(means):
            ax.annotate(f"{y:.2f}%", (x, y), xytext=(0, 14), textcoords="offset points", ha="center", weight="bold", color=TEAL)
        ax.set(xticks=[0, 1], xticklabels=["Random mean", "Real-Anchored TopK"], xlim=(-.25, 1.3), ylim=(60, 85), ylabel=f"Dev2000 {label} (%)")
        ax.set_title(f"{label}: +{means[1]-means[0]:.2f} pp", pad=16)
        ax.grid(axis="y", alpha=.15)
    fig.suptitle("Real-anchored selection improves the paired development results", fontsize=18, weight="bold", y=.98)
    savefig(fig, "01_TopK配对提升", "Legacy candidate pool: 640; train budget: 320. Each gray line is one model seed.\nRandom is averaged over five selection seeds within each model seed. Descriptive; final holdout not used.")
    data["TopK_Paired"] = (["ModelSeed", "Random_Accuracy_pct", "TopK_Accuracy_pct", "Random_MacroF1_pct", "TopK_MacroF1_pct"],
                            [[s, *panels[0][i], *panels[1][i]] for i, s in enumerate(seeds)])

    fig, ax = plt.subplots(figsize=(12, 5.7))
    gains = np.array([r["Delta"] * 100 for r in classes])
    bars = ax.bar([r["Class"] for r in classes], gains, color=[GRAY, GRAY, GRAY, TEAL], width=.58)
    ax.bar_label(bars, labels=[f"+{v:.2f}" for v in gains], padding=8, fontsize=16, weight="bold")
    ax.set(ylabel="TopK minus Random: F1 gain (percentage points)", ylim=(0, 13))
    ax.set_title("The largest mean gain is in Sci/Tech", fontsize=19, pad=22)
    ax.grid(axis="y", alpha=.15)
    savefig(fig, "02_类别收益", "Same legacy-pool Dev2000 protocol as Figure 01. Bars are mean differences, not confidence intervals.\nClass-level results are descriptive and do not establish a universal class-specific effect.")
    data["Class_Gains"] = (["Class", "Random_F1_pct", "TopK_F1_pct", "Delta_pp"],
                            [[r["Class"], r["Random_Mean_F1"]*100, r["TopK_Mean_F1"]*100, r["Delta"]*100] for r in classes])

    order = ["real", "qwen", "llama", "mix"]
    matrix_lookup = {r["Train_Source"]: r for r in matrix_rows}
    matrix = np.array([[matrix_lookup[s][t]*100 for t in order] for s in order])
    fig, ax = plt.subplots(figsize=(12, 6.6))
    im = ax.imshow(matrix, cmap="YlGnBu", vmin=0, vmax=100, aspect="auto")
    for i in range(4):
        for j in range(4):
            ax.text(j, i, f"{matrix[i,j]:.2f}", ha="center", va="center", fontsize=19, weight="bold", color="white" if matrix[i,j]>65 else NAVY)
    ax.set(xticks=range(4), xticklabels=[s.title() for s in order], yticks=range(4), yticklabels=[s.title() for s in order], xlabel="Test source", ylabel="Train source")
    ax.set_title("Synthetic holdout scores do not measure real generalization", fontsize=18, pad=24)
    colorbar=fig.colorbar(im, ax=ax, fraction=.03, pad=.03)
    colorbar.ax.set_title("Acc. (%)", fontsize=10, pad=12)
    savefig(fig, "03_来源兼容性矩阵", "Mean of 3 model seeds per cell; 48 training runs. Real training: 200 examples; synthetic training: 1,600.\nTraining budgets differ: this matrix is a domain diagnostic, not a controlled ranking of generator quality.")
    data["Source_Matrix"] = (["Train_Source", "Real", "Qwen", "Llama", "Mix"], [[s.title(), *matrix[i]] for i,s in enumerate(order)])

    # The CaR workbook has a missing DSIR row and an 80-run LN aggregate.
    # Use the original 25-run LN workbook for DSIR/LN, and CaR only for RA-CaR.
    selector = []
    for r in ln:
        label = {"Random":"Random", "TopK Proximity":"TopK", "Original CC-DSIR":"DSIR", "LN-CC-DSIR":"LN-DSIR"}[r["Method"]]
        selector.append([label, r["Dev_Macro_F1_Mean"]*100, r["Dev_Macro_F1_Std"]*100, r["Num_Training_Runs"]])
    r = next(r for r in car if r["Method"] == "RA-CaR")
    selector.append(["RA-CaR", r["Macro_F1_Mean"]*100, r["Macro_F1_Std"]*100, r["Num_Training_Runs"]])
    selector.sort(key=lambda r: ["Random", "DSIR", "LN-DSIR", "RA-CaR", "TopK"].index(r[0]))
    fig, ax = plt.subplots(figsize=(12, 5.7))
    colors = [TEAL if r[0]=="TopK" else GRAY for r in selector]
    ax.bar(range(5), [r[1] for r in selector], color=colors, alpha=.9, width=.58)
    ax.errorbar(range(5), [r[1] for r in selector], yerr=[r[2] for r in selector], fmt="none", color=NAVY, capsize=5, lw=1.4)
    for i,r in enumerate(selector):
        ax.text(i, r[1]+r[2]+1, f"{r[1]:.2f}%", ha="center", weight="bold")
    ax.set(xticks=range(5), xticklabels=[f"{r[0]}\n{r[3]} runs" for r in selector], ylabel="Dev2000 Macro-F1 (%)", ylim=(0,90))
    ax.set_title("More complex selectors did not exceed TopK in the legacy protocol", fontsize=18, pad=22)
    ax.grid(axis="y", alpha=.15)
    savefig(fig, "04_筛选器对照", "Mean +/- sample SD, not a confidence interval. TopK: 5 model runs; other arms: 5 selections x 5 model runs.\nDSIR/LN: original LN workbook (25 runs each); RA-CaR: CaR workbook. Different variance components enter these SDs.")
    data["Selector_Comparison"] = (["Method", "MacroF1_pct", "SD_pct", "N_Runs"], selector)

    fig, ax = plt.subplots(figsize=(12, 5.7))
    x = np.arange(len(ess)); old = np.array([r["Original_ESS"] for r in ess]); new = np.array([r["LN_ESS"] for r in ess])
    ax.bar(x-.18, old, width=.36, color=AMBER, label="Original DSIR")
    ax.bar(x+.18, new, width=.36, color=TEAL, label="Length-normalized DSIR")
    ax.set(xticks=x, xticklabels=[f'{r["Source"].title()}\n{r["Class"]}' for r in ess], ylabel="Effective sample size (pool size = 80)", ylim=(0,82))
    ax.set_title(f"Length normalization raises median ESS: {np.median(old):.2f} to {np.median(new):.2f}", fontsize=18, pad=22)
    ax.legend(frameon=False, loc="upper left"); ax.grid(axis="y", alpha=.15)
    savefig(fig, "05_DSIR有效样本量", "Eight source x class groups. ESS describes sampling-weight concentration, not downstream accuracy.\nThe improved diagnostic did not produce a mean Macro-F1 above TopK; see Figure 04.")
    data["DSIR_ESS"] = (["Group", "Original_ESS", "LN_ESS"], [[f'{r["Source"]}/{r["Class"]}',r["Original_ESS"],r["LN_ESS"]] for r in ess])

    vlookup = {(r["Method"],r["Metric"]):r for r in verifier}
    names = ["RandomMean", "TopK320", "SelfVerifier320", "CrossVerifier320"]
    vrows = [[n, vlookup[(n,"Dev_Macro_F1")]["Mean"]*100, vlookup[(n,"Dev_Macro_F1")]["Std"]*100,5] for n in names]
    fig, ax = plt.subplots(figsize=(12, 5.7))
    ax.bar(range(4),[r[1] for r in vrows],color=[GRAY,TEAL,GRAY,AMBER],width=.58)
    ax.errorbar(range(4),[r[1] for r in vrows],yerr=[r[2] for r in vrows],fmt="none",color=NAVY,capsize=6)
    for i,r in enumerate(vrows):ax.text(i,r[1]+r[2]+1,f"{r[1]:.2f}%",ha="center",weight="bold")
    ax.set(xticks=range(4),xticklabels=["Random mean","TopK","Self-verifier","Cross-verifier"],ylabel="Dev2000 Macro-F1 (%)",ylim=(0,92))
    ax.set_title("Cross-verification improves over self, but does not exceed TopK",fontsize=18,pad=22)
    ax.text(.5,.96,"Cross - Self: +2.02 pp (4/5 positive)    |    Cross - TopK: -0.65 pp (1/5 positive)",transform=ax.transAxes,ha="center",fontsize=12)
    ax.grid(axis="y",alpha=.15)
    savefig(fig,"06_Verifier下游对照","4,000-candidate pool; train budget: 320. Mean +/- SD across 5 model seeds; not a confidence interval.\nRandom is first averaged over 5 selections within each model seed. Do not mix with the legacy-pool figures.")
    data["Verifier_Comparison"] = (["Method","MacroF1_pct","SD_pct","N_ModelSeeds"],vrows)

    for name,(headers,rows) in data.items():
        with (TABLES/(name+".csv")).open("w",encoding="utf-8-sig",newline="") as f:
            writer=csv.writer(f);writer.writerow(headers);writer.writerows(rows)
    (TABLES/"chart_data.json").write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding="utf-8")
    build_workbook(data,manifest)
    build_slides()
    (OUT/"结果来源与核对记录.json").write_text(json.dumps({"sources":manifest,"notes":[
        "CaR Main_Comparison has Original DSIR N=0 and LN-DSIR N=80. Figure 04 takes DSIR/LN from the original 25-run LN workbook; CaR only supplies RA-CaR.",
        "No final-holdout evaluation or model training was run while building these materials.",
        "All figures are descriptive. No statistical significance claim is added.",
    ]},ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"source_workbooks":len(manifest),"figures":6,"slides":7,"output":str(OUT)},ensure_ascii=False))


def build_workbook(data, manifest):
    w=xlsxwriter.Workbook(OUT/"展示数据与可编辑图表.xlsx")
    header=w.add_format({"bold":True,"bg_color":NAVY,"font_color":"white","text_wrap":True})
    num=w.add_format({"num_format":"0.00"})
    note=w.add_format({"text_wrap":True,"valign":"top"})
    index=w.add_worksheet("阅读说明")
    notes=[
        ["用途","对应六张 PPT 图，包含可编辑数值和 Excel 原生图表。百分数数据已乘 100；差值单位为百分点。"],
        ["主线","01 TopK 配对提升；02 类别收益；03 来源矩阵；06 verifier。04、05 可作机制分析备份页。"],
        ["误差线","图 04/06 为标准差，不是置信区间。图 04 的 SD 包含不同随机性来源；不以误差线判断显著性。"],
        ["协议","图 01/02/04/05 使用早期候选池；图 06 使用 4,000 候选池并重新训练。数值不可跨阶段直接混排。"],
        ["来源矩阵","Real 训练 200 条，synthetic 训练 1,600 条；每单元 3 seeds。用于诊断，不是等预算来源排名。"],
        ["数据核对","CaR 总表中原始 DSIR 缺失、LN-DSIR 为 80 次汇总。因此图 04 的 DSIR/LN 使用独立 LN 原表的 25 次结果。"],
        ["原始文件","关键Excel/ 为六份原始工作簿的字节一致副本；不要直接向教授发送全部逐次结果，优先用本展示表和图。"],
        ["生成脚本","tools/build_research_presentation.py；依赖见 requirements-presentation.txt。"],
    ]
    index.set_column(0,0,18);index.set_column(1,1,115)
    for i,row in enumerate(notes):index.write_row(i,0,row,note);index.set_row(i,42)
    src=w.add_worksheet("来源索引");src.set_column(0,3,55)
    src.write_row(0,0,["key","copy","source","sha256"],header)
    for i,r in enumerate(manifest,1):src.write_row(i,0,[r[k] for k in ("key","copy","source","sha256")])
    for name,(headers,rows) in data.items():
        s=w.add_worksheet(name);s.freeze_panes(1,1);s.write_row(0,0,headers,header)
        s.set_column(0,0,25);s.set_column(1,len(headers)-1,22,num)
        for i,r in enumerate(rows,1):s.write_row(i,0,r)
        s.autofilter(0,0,len(rows),len(headers)-1)
        if name=="Source_Matrix":
            s.conditional_format(1,1,4,4,{"type":"3_color_scale","min_type":"num","min_value":0,"max_type":"num","max_value":100})
            s.insert_image("A9",str(FIG/"03_来源兼容性矩阵.png"),{"x_scale":.22,"y_scale":.22})
            continue
        if name=="TopK_Paired":
            for metric,cols,anchor in [("Accuracy",(1,2),"G2"),("Macro-F1",(3,4),"G20")]:
                chart=w.add_chart({"type":"column"})
                for col in cols:chart.add_series({"name":[name,0,col],"categories":[name,1,0,5,0],"values":[name,1,col,5,col]})
                chart.set_title({"name":f"Paired model seeds: {metric}"});chart.set_y_axis({"name":"Percent","min":0,"max":100});chart.set_style(10);s.insert_chart(anchor,chart)
            continue
        chart=w.add_chart({"type":"column"})
        if name=="DSIR_ESS":cols=[1,2]
        elif name=="Class_Gains":cols=[3]
        else:cols=[1]
        for col in cols:
            series_color=AMBER if name=="DSIR_ESS" and col==1 else TEAL
            spec={"name":[name,0,col],"categories":[name,1,0,len(rows),0],"values":[name,1,col,len(rows),col],"fill":{"color":series_color},"data_labels":{"value":True,"num_format":"0.00"}}
            if name in ("Selector_Comparison","Verifier_Comparison"):
                spec["y_error_bars"]={"type":"custom","plus_values":[name,1,2,len(rows),2],"minus_values":[name,1,2,len(rows),2]}
            chart.add_series(spec)
        unit="ESS" if name=="DSIR_ESS" else ("F1 gain (percentage points)" if name=="Class_Gains" else "Macro-F1 (%)")
        chart.set_title({"name":name.replace("_"," ")});chart.set_y_axis({"min":0,"name":unit});chart.set_style(10);chart.set_size({"width":820,"height":440});s.insert_chart("G2",chart)
    w.close()


def text(slide,x,y,width,height,content,size=20,color=NAVY,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(width),Inches(height))
    tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(content.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name="Microsoft YaHei";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip("#"))
    return box


def build_slides():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    s=prs.slides.add_slide(prs.slide_layouts[6])
    text(s,.7,.7,12,1.2,"固定预算下的合成文本筛选\n与真实域泛化",32,bold=True)
    text(s,.75,2.3,11.8,2.5,"AG News 四分类 · BERT 下游分类器\n研究问题：哪些合成样本值得进入有限的训练预算？\n从真实数据锚定筛选，到分布诊断与 verifier 的效用边界。",23)
    text(s,.75,5.55,11.8,1.2,"讨论用图表初稿 · 基于已有结果整理\n姓名、申请方向与学校信息可自行补充。数值为开发或诊断结果，最终 holdout 未使用。",15,color="#53697D")
    entries=[
        ("01_TopK配对提升","真实数据锚定筛选：配对开发集结果更好","Accuracy +2.67 pp；Macro-F1 +3.81 pp；五个配对模型 seed 均正向。","每个 Random 点为同一模型 seed 下五次选样的均值；单一早期候选池。"),
        ("02_类别收益","类别诊断：Sci/Tech 的平均收益最大","Sci/Tech F1 平均提高 10.13 个百分点；其他类别收益较小。","均值差异为描述性结果，不解释为普遍的类别效应。"),
        ("03_来源兼容性矩阵","来源兼容性：合成 holdout 高分不代表真实泛化","合成训练在合成测试上约 97%–100%；在真实测试上约 68%–70%。","Real 训练 200 条、synthetic 训练 1,600 条；三个 seed/单元，不能作等预算来源排名。"),
        ("04_筛选器对照","机制对照：复杂筛选器未超过 TopK","保留 DSIR、LN-DSIR、RA-CaR 的结果，用于分析方法何时有用。","误差线为 SD。DSIR/LN 来自独立 25 次原表，避开 CaR 总表中的缺失与混合汇总。"),
        ("05_DSIR有效样本量","权重诊断：修复 ESS 集中仍不足以提升效用","长度归一化使 ESS 中位数从约 1.00 提升到 49.58。","ESS 衡量采样权重集中程度；它的改善不等于下游泛化提升。"),
        ("06_Verifier下游对照","Verifier：cross 优于 self，仍未胜过 TopK","Cross − Self：+2.02 pp（4/5 正向）；Cross − TopK：−0.65 pp（1/5 正向）。","4,000 候选池独立阶段；误差线为五个模型 seed 的 SD。按预定规则归类为 No useful signal。"),
    ]
    for i,(fname,title,claim,foot) in enumerate(entries,1):
        s=prs.slides.add_slide(prs.slide_layouts[6]);text(s,.5,.25,12.3,.5,title,25,bold=True)
        image_path=FIG/(fname+".png")
        with Image.open(image_path) as img:
            aspect=img.width/img.height
        height=min(5.25,11.7/aspect);width=height*aspect
        s.shapes.add_picture(str(image_path),Inches((13.333-width)/2),Inches(.95+(5.25-height)/2),width=Inches(width),height=Inches(height))
        text(s,.65,6.2,12,.5,claim,18,color=TEAL,bold=True)
        text(s,.65,6.78,11.8,.5,foot,11,color="#53697D")
        s.notes_slide.notes_text_frame.text=f"来源和作图数据见 展示数据与可编辑图表.xlsx。图片：{fname}.png。\n{foot}\n展示图为 PNG；修改数值请在 Excel 原生图表或生成脚本中完成。"
        text(s,12.4,7.1,.6,.25,f"{i:02d}",11,color=GRAY)
    prs.save(OUT/"申请交流图表初稿.pptx")


if __name__=="__main__":main()
